"""Readers for exposure denominators: the driver census and the ITV kilometre estimates."""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pandas as pd

from dgt_stats import agebands, io_population
from dgt_stats.paths import (
    CENSUS_AGE_YEARS,
    CENSUS_TABLES_2025_PATH,
    CENSUS_YEARS,
    INTERIM_DATA_DIR,
    KM_BY_OWNER_AGE_2024_PATH,
    KM_ESTIMATED_2022_PATH,
    KM_MEAN_2022_PATH,
    KM_MEAN_BY_TYPE_2024_PATH,
    census_age_raw_path,
    census_raw_path,
    census_tables_raw_path,
)

log = logging.getLogger(__name__)

SEX_CODES = {"V": "male", "M": "female"}

# Driver census by age: text-file encodings and the published class × age sheets by sex.
CENSUS_AGE_ENCODINGS = {2023: "latin-1", 2024: "latin-1", 2025: "utf-8-sig"}
CENSUS_AGE_TABLE_YEARS = tuple(range(2014, 2024))
CENSUS_AGE_SHEETS = {"P_6_1_1_7": "total", "P_6_1_2_7": "male", "P_6_1_3_7": "female"}
CENSUS_TOTAL_ROWS = {"TOTAL GENERAL", "TOTAL CENSO"}

KM_VEHICLE_TYPES = (
    "Ciclomotor",
    "Motocicleta",
    "Turismo",
    "Furgoneta",
    "Camión hasta 3.500Kg",
    "Camión más de 3.500Kg",
    "Autobús",
)


def read_census_year(year: int) -> pd.DataFrame:
    """One census extract: drivers by province, sex, highest licence class and year of issue.

    ``CLASE_PERMISO`` is the driver's single highest class, so the file sums to the number of
    licensed drivers, not to the number of permits.
    """
    raw = pd.read_csv(census_raw_path(year), sep="|", dtype=str, encoding="utf-8-sig")
    raw.columns = [column.strip() for column in raw.columns]
    for column in raw.columns:
        raw[column] = raw[column].str.strip()
    out = pd.DataFrame(
        {
            "census_year": year,
            "province_code": raw["COD_PROVINCIA"],
            "sex_code": raw["IND_SEXO"],
            "sex": raw["IND_SEXO"].map(SEX_CODES),
            "licence_class": raw["CLASE_PERMISO"],
            "licence_year": raw["DESC_ANTIG_PERMISO"],
            "n_drivers": pd.to_numeric(raw["NUM_CONDUCTORES"]).astype("int32"),
        }
    )
    if out["sex"].isna().any():
        raise ValueError(f"census {year}: unexpected sex code")
    return out.astype(
        {
            "census_year": "int16",
            "province_code": "string",
            "sex_code": "string",
            "sex": "string",
            "licence_class": "string",
            "licence_year": "string",
        }
    )


def read_census_all(years: tuple[int, ...] = CENSUS_YEARS) -> pd.DataFrame:
    return pd.concat([read_census_year(year) for year in years], ignore_index=True)


def _fine_band(low: int | None, high: int | None) -> str:
    key = agebands.band_for(low, high, agebands.DGT_BANDS)
    if key is None:
        raise ValueError(f"age interval {low}-{high} is outside the driver bands")
    return key


def read_census_age_year(year: int) -> pd.DataFrame:
    """Driver census by province × sex × age band for 2023–2025 (pipe-delimited text).

    ``n_drivers`` is the full census (people holding a driving permit or a moped / agricultural
    licence, counted once), which equals the published census total; ``n_permit_holders`` counts
    permit holders only and ``n_licence_only`` licence holders, some of whom also hold a permit.
    """
    raw = pd.read_csv(
        census_age_raw_path(year), sep="|", dtype=str, encoding=CENSUS_AGE_ENCODINGS[year]
    )
    raw.columns = [column.strip() for column in raw.columns]
    for column in raw.columns:
        raw[column] = raw[column].str.strip()
    parsed = [agebands.parse_age_label(label) for label in raw["EDAD"]]
    out = pd.DataFrame(
        {
            "census_year": year,
            "province_code": raw["COD_PROVINCIA"].str.zfill(2),
            "sex": raw["IND_SEXO"].map(SEX_CODES),
            "age_label": raw["EDAD"],
            "age_low": pd.array(
                [None if item is None else item[0] for item in parsed], dtype="Int16"
            ),
            "age_high": pd.array(
                [None if item is None else item[1] for item in parsed], dtype="Int16"
            ),
            "band": [agebands.UNKNOWN if item is None else _fine_band(*item) for item in parsed],
            "n_permit_holders": pd.to_numeric(raw["NUM_PERMISOS"]).astype("int64"),
            "n_licence_only": pd.to_numeric(raw["NUM_LICENCIAS"]).astype("int64"),
            "n_drivers": pd.to_numeric(raw["NUM_LICENCIAS_PERMISOS"]).astype("int64"),
        }
    )
    if out.sex.isna().any():
        raise ValueError(f"census by age {year}: unexpected sex code")
    lower = out[["n_permit_holders", "n_licence_only"]].max(axis=1)
    upper = out.n_permit_holders + out.n_licence_only
    if ((out.n_drivers < lower) | (out.n_drivers > upper)).any():
        raise ValueError(f"census by age {year}: total outside [max(permits, licences), sum]")
    return out.astype(
        {
            "census_year": "int16",
            "province_code": "string",
            "sex": "string",
            "age_label": "string",
            "band": "string",
        }
    )


def read_census_age_all(years: tuple[int, ...] = CENSUS_AGE_YEARS) -> pd.DataFrame:
    return pd.concat([read_census_age_year(year) for year in years], ignore_index=True)


def read_census_age_tables(year: int) -> pd.DataFrame:
    """Published census by licence class × age band (sheets P.6.1.1.7 and the men/women variants).

    Only the census total row is kept (``TOTAL GENERAL`` up to 2020, ``Total censo`` from 2021): one
    row per sex × age band with the number of drivers. The parsed bands must sum to the published
    total column, otherwise the reader raises.
    """
    workbook = openpyxl.load_workbook(census_tables_raw_path(year), read_only=True, data_only=True)
    records: list[dict[str, object]] = []
    try:
        for sheet, sex in CENSUS_AGE_SHEETS.items():
            rows = [tuple(row) for row in workbook[sheet].iter_rows(values_only=True)]
            texts = [
                [" ".join(str(cell).split()) if cell is not None else "" for cell in row]
                for row in rows
            ]
            header_index = next(
                i for i, row in enumerate(texts) if any(cell.startswith("15 a 17") for cell in row)
            )
            total_index = max(
                i for i, row in enumerate(texts) if row and row[0].upper() in CENSUS_TOTAL_ROWS
            )
            header, total = texts[header_index], rows[total_index]
            published_total = None
            parsed_sum = 0
            for column, label in enumerate(header):
                if label.lower().startswith("total"):
                    published_total = float(total[column])
                    continue
                if not label or label.upper() == "CLASES":
                    continue
                parsed = agebands.parse_age_label(label)
                value = int(total[column] or 0)
                parsed_sum += value
                records.append(
                    {
                        "census_year": year,
                        "sex": sex,
                        "age_label": label,
                        "age_low": None if parsed is None else parsed[0],
                        "age_high": None if parsed is None else parsed[1],
                        "band": agebands.UNKNOWN if parsed is None else _fine_band(*parsed),
                        "n_drivers": value,
                        "source_sheet": sheet,
                    }
                )
            if published_total is None or parsed_sum != published_total:
                raise ValueError(
                    f"census tables {year} {sheet}: bands sum to {parsed_sum}, "
                    f"published total {published_total}"
                )
    finally:
        workbook.close()
    out = pd.DataFrame.from_records(records)
    return out.astype(
        {
            "census_year": "int16",
            "sex": "string",
            "age_label": "string",
            "age_low": "Int16",
            "age_high": "Int16",
            "band": "string",
            "n_drivers": "int64",
            "source_sheet": "string",
        }
    )


def read_census_age_tables_all(
    years: tuple[int, ...] = CENSUS_AGE_TABLE_YEARS,
) -> pd.DataFrame:
    return pd.concat([read_census_age_tables(year) for year in years], ignore_index=True)


def licence_holders_by_age() -> pd.DataFrame:
    """Licensed drivers by year (2014–2025), sex (male, female, total) and fine age band, national.

    2014–2023 come from the published class × age tables (one publication type for the whole
    series), 2024–2025 from the text files summed over provinces; ``source`` says which. The 2023
    text file differs from the 2023 tables by up to about 1 % in some bands (different extraction
    dates); the validation layer records that gap.
    """
    tables = read_census_age_tables_all()
    from_tables = (
        tables.groupby(["census_year", "sex", "band"], observed=True)
        .n_drivers.sum()
        .reset_index()
        .assign(source="census_tables")
    )
    text = read_census_age_all()
    text = text[text.census_year > max(CENSUS_AGE_TABLE_YEARS)]
    by_sex = (
        text.groupby(["census_year", "sex", "band"], observed=True).n_drivers.sum().reset_index()
    )
    total = (
        text.groupby(["census_year", "band"], observed=True)
        .n_drivers.sum()
        .reset_index()
        .assign(sex="total")
    )
    from_text = pd.concat([by_sex, total], ignore_index=True).assign(source="census_text")
    out = pd.concat([from_tables, from_text], ignore_index=True).rename(
        columns={"census_year": "year"}
    )
    order = {key: i for i, key in enumerate([*agebands.DGT_BANDS, agebands.UNKNOWN])}
    out["band_order"] = out.band.map(order)
    out = out.sort_values(["year", "sex", "band_order"]).drop(columns="band_order")
    return out.reset_index(drop=True).astype(
        {"year": "int16", "sex": "string", "band": "string", "source": "string"}
    )


def read_census_province_totals_2025() -> pd.DataFrame:
    """Published 2025 census by province of residence and sex (sheet P_6_1_1_10)."""
    workbook = openpyxl.load_workbook(CENSUS_TABLES_2025_PATH, read_only=True, data_only=True)
    rows = [tuple(row) for row in workbook["P_6_1_1_10"].iter_rows(values_only=True)]
    workbook.close()
    header = next(i for i, row in enumerate(rows) if row and str(row[0]).strip() == "PROVINCIAS")
    records = []
    for row in rows[header + 1 :]:
        if not row or row[0] is None:
            continue
        name = " ".join(str(row[0]).split())
        records.append(
            {
                "province": name,
                "is_total": name.lower() == "total",
                "drivers_excluding_licences": row[3],
                "licence_only": row[6],
                "men": row[7],
                "women": row[8],
                "total": row[9],
                "source_sheet": "P_6_1_1_10",
            }
        )
    out = pd.DataFrame.from_records(records)
    for column in ("drivers_excluding_licences", "licence_only", "men", "women", "total"):
        out[column] = pd.to_numeric(out[column]).astype("int64")
    return out.astype({"province": "string", "source_sheet": "string"})


def read_km_mean_2022() -> pd.DataFrame:
    """Fleet size and mean annual km by vehicle type and age band for 2022, plus vehicle-km."""
    workbook = openpyxl.load_workbook(KM_MEAN_2022_PATH, read_only=True, data_only=True)
    rows = [tuple(row) for row in workbook["Hoja1"].iter_rows(values_only=True)]
    workbook.close()
    band_row = next(i for i, row in enumerate(rows) if row and row[1] == "De 0 a 4 años")
    bands = [cell for cell in rows[band_row][1:] if cell is not None]
    records = []
    for row in rows[band_row + 1 :]:
        if not row or row[0] not in KM_VEHICLE_TYPES:
            continue
        for index, band in enumerate(bands):
            n_vehicles = row[1 + 2 * index]
            mean_km = row[2 + 2 * index]
            records.append(
                {
                    "year": 2022,
                    "vehicle_type": row[0],
                    "age_band": band,
                    "n_vehicles": int(n_vehicles),
                    "mean_km_year": float(mean_km),
                    "vehicle_km": int(n_vehicles) * float(mean_km),
                    "source_sheet": "Hoja1",
                }
            )
    out = pd.DataFrame.from_records(records)
    if len(out) != len(KM_VEHICLE_TYPES) * len(bands):
        raise ValueError("km mean table: unexpected number of strata")
    return out.astype(
        {"year": "int16", "vehicle_type": "string", "age_band": "string", "source_sheet": "string"}
    )


def read_km_estimated_2022() -> pd.DataFrame:
    """Estimated annual km per vehicle by stratum, all six vehicle-type sheets stacked."""
    workbook = openpyxl.load_workbook(KM_ESTIMATED_2022_PATH, read_only=True, data_only=True)
    frames = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() for cell in next(rows)]
        records = [row for row in rows if any(cell is not None for cell in row)]
        frame = pd.DataFrame.from_records(records, columns=header)
        frame["source_sheet"] = sheet.title
        frames.append(frame)
    workbook.close()
    out = pd.concat(frames, ignore_index=True)
    out = out.rename(
        columns={
            "FECHA PARQUE": "year",
            "TIPO_VEHICULO": "vehicle_type",
            "EMISIONES_NORM": "emission_standard",
            "ANTIGUEDAD": "age_band",
            "CILINDRADA": "engine_size",
            "CARGA": "payload",
            "DESC_PROPULSION": "fuel",
            "kmRecorridos": "km_year",
        }
    )
    out["year"] = pd.to_numeric(out["year"]).astype("int16")
    out["km_year"] = pd.to_numeric(out["km_year"]).astype("float64")
    for column in (
        "vehicle_type",
        "emission_standard",
        "age_band",
        "engine_size",
        "payload",
        "fuel",
    ):
        out[column] = out[column].astype("string").str.strip()
    return out


# DGT's 2024 kilometre release groups the fleet by the owner's age. The bands are its own; the
# company row has no age and is kept as a row of its own rather than being spread over the others.
KM_OWNER_CATEGORIES = {
    "Ciclomotores": "moped",
    "Motocicletas": "motorcycle",
    "Turismos": "car",
    "Furgonetas": "van",
    "Camiones (hasta 3.500Kg MMA)": "light_truck",
    "Camiones (desde 3.500Kg MMA)": "heavy_truck",
    "Tractores industriales": "tractor_unit",
    "Autobuses": "bus",
}
KM_OWNER_COMPANY = "Vehículo a nombre de empresa"
KM_OWNER_TOTALS = {
    "TURISMOS": "car",
    "FURGONETAS": "van",
    "CAMIONES <= 3500KG": "light_truck",
    "CICLOMOTORES": "moped",
    "CAMIONES > 3500KG": "heavy_truck",
    "MOTOCICLETAS": "motorcycle",
    "TRACTORES INDUSTRIALES": "tractor_unit",
    "AUTOBUSES": "bus",
}
# The owner-age table leaves out the vehicles whose owner's age DGT could not classify; the check
# below allows for that and no more.
KM_OWNER_TOLERANCE = 0.005


def read_km_by_owner_age_2024() -> pd.DataFrame:
    """Vehicles and annual kilometres by vehicle category and the owner's age band, 2024.

    DGT's *Kilómetros anualizados recorridos por el parque móvil 2024* publishes this table in its
    additional material. It is the only Spanish source that puts distance driven and a person's age
    in the same cell, and it does it for the whole circulating fleet rather than for a survey
    sample. What it is not is the *driver's* age: a car registered to a person of 75 may be driven
    by someone else, and the vehicles registered to companies carry no age at all. The reader keeps
    the company row so that the share of kilometres it holds is visible wherever the table is used.

    The parsed rows are reconciled against the same release's table 6 (vehicles and mean annual km
    by category) within 0.5 %, the margin left by the owners DGT could not classify.
    """
    detail = pd.read_excel(KM_BY_OWNER_AGE_2024_PATH, sheet_name="Detalle 2024", header=0)
    detail.columns = ["category_es", "owner_band", "n_vehicles", "total_km", "mean_km"]
    unknown = set(detail.category_es.unique()) - set(KM_OWNER_CATEGORIES)
    if unknown:
        raise ValueError(f"km by owner age: unexpected vehicle categories {sorted(unknown)}")
    out = pd.DataFrame(
        {
            "year": 2024,
            "vehicle_group": detail.category_es.map(KM_OWNER_CATEGORIES),
            "category_es": detail.category_es,
            "owner_band": detail.owner_band.astype(str).str.strip(),
            "n_vehicles": pd.to_numeric(detail.n_vehicles).astype("int64"),
            "total_km": pd.to_numeric(detail.total_km).astype("int64"),
            "mean_km": pd.to_numeric(detail.mean_km).astype("float64"),
        }
    )
    out["is_company"] = out.owner_band == KM_OWNER_COMPANY
    out["band"] = [
        None
        if company
        else agebands.band_for(
            *agebands.parse_age_label(_owner_label(label)), agebands.EXPOSURE_BANDS
        )
        for label, company in zip(out.owner_band, out.is_company)
    ]
    published = read_km_means_2024().set_index("vehicle_group")
    parsed = out.groupby("vehicle_group")[["n_vehicles", "total_km"]].sum()
    for group, row in parsed.iterrows():
        for column in ("n_vehicles", "total_km"):
            reference = float(published.loc[group, column])
            if abs(row[column] - reference) / reference > KM_OWNER_TOLERANCE:
                raise ValueError(
                    f"km by owner age: {group} {column} {row[column]:,.0f} is more than "
                    f"{KM_OWNER_TOLERANCE:.1%} from the published {reference:,.0f}"
                )
    return out.astype(
        {
            "year": "int16",
            "vehicle_group": "string",
            "category_es": "string",
            "owner_band": "string",
            "band": "string",
        }
    )


def _owner_label(label: str) -> str:
    """DGT's owner bands as an age label ``parse_age_label`` understands."""
    text = str(label).strip()
    if text.endswith("+"):
        return f"{text[:-1]} y más"
    if "-" in text:
        low, high = text.split("-", 1)
        return f"de {low} a {high}"
    return f"de {text} a {text}"


def read_km_means_2024() -> pd.DataFrame:
    """Vehicles, total and mean annual kilometres by vehicle category, 2024 (release table 6)."""
    frame = pd.read_excel(KM_MEAN_BY_TYPE_2024_PATH, sheet_name="Detalle 2024", header=0)
    frame.columns = ["category_es", "n_vehicles", "total_km", "mean_km"]
    unknown = set(frame.category_es.unique()) - set(KM_OWNER_TOTALS)
    if unknown:
        raise ValueError(f"km means 2024: unexpected categories {sorted(unknown)}")
    out = pd.DataFrame(
        {
            "year": 2024,
            "vehicle_group": frame.category_es.map(KM_OWNER_TOTALS),
            "category_es": frame.category_es,
            "n_vehicles": pd.to_numeric(frame.n_vehicles).astype("int64"),
            "total_km": pd.to_numeric(frame.total_km).astype("int64"),
            "mean_km": pd.to_numeric(frame.mean_km).astype("float64"),
        }
    )
    return out.astype({"year": "int16", "vehicle_group": "string", "category_es": "string"})


EXPOSURE_BUILDERS = {
    "censo_conductores": read_census_all,
    "censo_provincias_2025": read_census_province_totals_2025,
    "censo_edad": read_census_age_all,
    "censo_edad_tablas": read_census_age_tables_all,
    "conductores_por_edad": licence_holders_by_age,
    "poblacion_ine": io_population.read_population,
    "km_medios_2022": read_km_mean_2022,
    "km_estimados_2022": read_km_estimated_2022,
    "km_edad_propietario_2024": read_km_by_owner_age_2024,
    "km_medios_tipo_2024": read_km_means_2024,
}


def interim_path(name: str) -> Path:
    return INTERIM_DATA_DIR / f"{name}.parquet"


def build_exposure(force: bool = False) -> list[Path]:
    written: list[Path] = []
    INTERIM_DATA_DIR.mkdir(parents=True, exist_ok=True)
    for name, builder in EXPOSURE_BUILDERS.items():
        target = interim_path(name)
        if target.exists() and not force:
            log.info("exposure %s: exists, skipping", name)
            written.append(target)
            continue
        frame = builder()
        frame.to_parquet(target, index=False)
        log.info("exposure %s: %s rows -> %s", name, f"{len(frame):,}", target.name)
        written.append(target)
    return written


def read_exposure(name: str) -> pd.DataFrame:
    return pd.read_parquet(interim_path(name))
