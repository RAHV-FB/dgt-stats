"""Readers for the published DGT tables: the 1993–2024 series and the yearly statistical tables.

Every function returns a tidy ``pandas.DataFrame`` and records the sheet it came from in a
``source_sheet`` column. Sheet layouts are located by the label in the first header cell rather than
by fixed row numbers so that small layout shifts in future releases are caught as parse errors.
"""

from __future__ import annotations

import logging
from pathlib import Path

import openpyxl
import pandas as pd

from dgt_stats import agebands
from dgt_stats.paths import (
    INTERIM_DATA_DIR,
    SERIES_PATH,
    TABLE_YEARS,
    TABLES_2024_PATH,
    tables_raw_path,
)

log = logging.getLogger(__name__)

MONTH_ABBREVIATIONS = (
    "Ene.",
    "Feb.",
    "Mar.",
    "Abr.",
    "May.",
    "Jun.",
    "Jul.",
    "Ago.",
    "Sep.",
    "Oct.",
    "Nov.",
    "Dic.",
)
MONTH_NAMES = (
    "Enero",
    "Febrero",
    "Marzo",
    "Abril",
    "Mayo",
    "Junio",
    "Julio",
    "Agosto",
    "Septiembre",
    "Octubre",
    "Noviembre",
    "Diciembre",
)

# Monthly sheets in the series workbook: sheet -> (metric, zone).
MONTHLY_SHEETS: dict[str, tuple[str, str]] = {
    "M_I-U_Meses": ("deaths_30d", "all"),
    "M_I_Meses": ("deaths_30d", "interurban"),
    "M_U_Meses": ("deaths_30d", "urban"),
    "HG_I-U_Meses": ("hospitalised_30d", "all"),
    "HG_I_Meses": ("hospitalised_30d", "interurban"),
    "HG_U_Meses": ("hospitalised_30d", "urban"),
    "HL_I-U_Meses": ("non_hospitalised_30d", "all"),
    "HL_I_Meses": ("non_hospitalised_30d", "interurban"),
    "HL_U_Meses": ("non_hospitalised_30d", "urban"),
    "Vict_I-U_Meses": ("victims_30d", "all"),
    "Vict_I_Meses": ("victims_30d", "interurban"),
    "Vict_U_Meses": ("victims_30d", "urban"),
    "Fallec_24h_I_U": ("deaths_24h", "all"),
    "Fallec_24h_I": ("deaths_24h", "interurban"),
    "Fallec_24h_U": ("deaths_24h", "urban"),
}

# Province sheets: sheet -> (metric, zone).
PROVINCE_SHEETS: dict[str, tuple[str, str]] = {
    "Acc_Vict_I-U_Provincias": ("crashes", "all"),
    "Acc_Vict_I_Provincias": ("crashes", "interurban"),
    "Acc_Vict_U_Provincias": ("crashes", "urban"),
    "MU_I-U_Provincias": ("deaths_30d", "all"),
    "MU_I_Provincias": ("deaths_30d", "interurban"),
    "MU_U_Provincias": ("deaths_30d", "urban"),
    "HG_I-U-Provincias": ("hospitalised_30d", "all"),
    "HG_I_Provincias": ("hospitalised_30d", "interurban"),
    "HG_U_Provincias": ("hospitalised_30d", "urban"),
    "HL_I-U_Provincias": ("non_hospitalised_30d", "all"),
    "HL_I_Provincias": ("non_hospitalised_30d", "interurban"),
    "HL_U_Provincias": ("non_hospitalised_30d", "urban"),
    "VIC_I-U_Provincias": ("victims_30d", "all"),
    "VIC_I_Provincias": ("victims_30d", "interurban"),
    "VIC_U_Provincias": ("victims_30d", "urban"),
}

# Driver / occupant death sheets by vehicle type: sheet -> (population, severity, zone).
ROAD_USER_SHEETS: dict[str, tuple[str, str, str]] = {
    "Cond-Pasj_MU_I-U": ("drivers_and_passengers", "deaths_30d", "all"),
    "Cond-Pasj_MU_I": ("drivers_and_passengers", "deaths_30d", "interurban"),
    "Cond-Pasj_MU_U": ("drivers_and_passengers", "deaths_30d", "urban"),
    "Cond-Pasj_HG_I-U": ("drivers_and_passengers", "hospitalised_30d", "all"),
    "Cond-Pasj_HG_I": ("drivers_and_passengers", "hospitalised_30d", "interurban"),
    "Cond-Pasj_HG_U": ("drivers_and_passengers", "hospitalised_30d", "urban"),
    "Cond-Pasj_HL_I-U": ("drivers_and_passengers", "non_hospitalised_30d", "all"),
    "Cond-Pasj_HL_I": ("drivers_and_passengers", "non_hospitalised_30d", "interurban"),
    "Cond-Pasj_HL_U": ("drivers_and_passengers", "non_hospitalised_30d", "urban"),
    "Cond_MU_I-U": ("drivers", "deaths_30d", "all"),
    "Cond_MU_I": ("drivers", "deaths_30d", "interurban"),
    "Cond_MU_U": ("drivers", "deaths_30d", "urban"),
    "Cond_HG_I-U": ("drivers", "hospitalised_30d", "all"),
    "Cond_HG_I": ("drivers", "hospitalised_30d", "interurban"),
    "Cond_HG_U": ("drivers", "hospitalised_30d", "urban"),
    "Cond_HL_I-U": ("drivers", "non_hospitalised_30d", "all"),
    "Cond_HL_I": ("drivers", "non_hospitalised_30d", "interurban"),
    "Cond_HL_U": ("drivers", "non_hospitalised_30d", "urban"),
}

AGE_SHEETS: dict[str, str] = {"Vict_ED_I_U": "all", "Vict_ED_I": "interurban", "Vict_ED_U": "urban"}
SEX_SHEETS: dict[str, str] = {"Vict_SX_I_U": "all", "Vict_SX_I": "interurban", "Vict_SX_U": "urban"}
PEDESTRIAN_SHEETS: dict[str, str] = {
    "Pea_Vic_I-U": "all",
    "Pea_Vic_I": "interurban",
    "Pea_Vic_U": "urban",
}
PEDESTRIAN_AGE_SHEETS: dict[str, str] = {
    "Edad_Pea_VIC_I-U": "all",
    "Edad_Pea_VIC_I": "interurban",
    "Edad_Pea_VIC_U": "urban",
}

SEVERITY_LEVELS = ("deaths_30d", "hospitalised_30d", "non_hospitalised_30d")
SEVERITY_LABELS = {
    "Fallecidos": "deaths_30d",
    "Heridos hospitalizados": "hospitalised_30d",
    "Heridos no hospitalizados": "non_hospitalised_30d",
    "Total": "victims_30d",
    "TOTAL": "victims_30d",
}


# --------------------------------------------------------------------------- generic helpers


_WORKBOOKS: dict[Path, openpyxl.Workbook] = {}


def _workbook(path: Path) -> openpyxl.Workbook:
    """Open a workbook once per process; ``close_workbooks`` releases the file handles."""
    if path not in _WORKBOOKS:
        _WORKBOOKS[path] = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return _WORKBOOKS[path]


def close_workbooks() -> None:
    """Close every cached read-only workbook and drop it from the cache."""
    for workbook in _WORKBOOKS.values():
        workbook.close()
    _WORKBOOKS.clear()


def _rows(path: Path, sheet: str) -> list[tuple]:
    return [tuple(row) for row in _workbook(path)[sheet].iter_rows(values_only=True)]


def _text(value: object) -> str:
    return "" if value is None else " ".join(str(value).split())


def _find_header(rows: list[tuple], first_cell: str) -> int:
    for index, row in enumerate(rows):
        if row and _text(row[0]).lower() == first_cell.lower():
            return index
    raise ValueError(f"header row starting with {first_cell!r} not found")


def _number(value: object) -> float | None:
    """Numeric cell or ``None``; the series workbook uses '.' for not-available."""
    if value is None or (isinstance(value, str) and value.strip() in {"", ".", "-"}):
        return None
    return float(value)


def _year_rows(rows: list[tuple], header_index: int) -> list[tuple]:
    body = []
    for row in rows[header_index + 1 :]:
        if not row or row[0] is None:
            continue
        try:
            year = int(row[0])
        except (TypeError, ValueError):
            continue
        body.append((year, *row[1:]))
    return body


# --------------------------------------------------------------------------- series workbook


def read_series_annual() -> pd.DataFrame:
    """Annual national totals 1993–2024: crashes by zone, victims by severity, and DGT rate table."""
    frames: list[pd.DataFrame] = []

    rows = _rows(SERIES_PATH, "Acc_Vict")
    header = _find_header(rows, "Años")
    for year, total, _, interurban, _, urban, *_ in _year_rows(rows, header):
        frames.append(
            pd.DataFrame(
                {
                    "year": year,
                    "metric": "crashes",
                    "zone": ["all", "interurban", "urban"],
                    "value": [_number(total), _number(interurban), _number(urban)],
                    "source_sheet": "Acc_Vict",
                }
            )
        )

    rows = _rows(SERIES_PATH, "Vict_I-U")
    header = _find_header(rows, "Años")
    for year, total, deaths, hosp, non_hosp, *_ in _year_rows(rows, header):
        frames.append(
            pd.DataFrame(
                {
                    "year": year,
                    "metric": [
                        "victims_30d",
                        "deaths_30d",
                        "hospitalised_30d",
                        "non_hospitalised_30d",
                    ],
                    "zone": "all",
                    "value": [_number(total), _number(deaths), _number(hosp), _number(non_hosp)],
                    "source_sheet": "Vict_I-U",
                }
            )
        )

    rows = _rows(SERIES_PATH, "Tasas_Acc_Vic")
    header = _find_header(rows, "Años")
    rate_metrics = (
        "vehicle_fleet",
        "crashes_per_10k_vehicles",
        "deaths_per_10k_vehicles",
        "deaths_per_1k_crashes",
        "injured_per_1k_crashes",
        "deaths_per_10k_population",
    )
    for year, *values in _year_rows(rows, header):
        frames.append(
            pd.DataFrame(
                {
                    "year": year,
                    "metric": list(rate_metrics),
                    "zone": "all",
                    "value": [_number(v) for v in values[: len(rate_metrics)]],
                    "source_sheet": "Tasas_Acc_Vic",
                }
            )
        )

    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {"year": "int16", "metric": "string", "zone": "string", "source_sheet": "string"}
    )


def read_series_monthly() -> pd.DataFrame:
    """Every monthly sheet stacked: year × month × metric × zone."""
    frames: list[pd.DataFrame] = []
    for sheet, (metric, zone) in MONTHLY_SHEETS.items():
        rows = _rows(SERIES_PATH, sheet)
        header = _find_header(rows, "Años")
        labels = [_text(cell) for cell in rows[header][1:13]]
        if tuple(labels) != MONTH_ABBREVIATIONS:
            raise ValueError(f"{sheet}: unexpected month header {labels}")
        for year, *values in _year_rows(rows, header):
            months = [_number(v) for v in values[:12]]
            total = _number(values[12]) if len(values) > 12 else None
            frames.append(
                pd.DataFrame(
                    {
                        "year": year,
                        "month": range(1, 13),
                        "metric": metric,
                        "zone": zone,
                        "value": months,
                        "annual_total": total,
                        "source_sheet": sheet,
                    }
                )
            )
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "year": "int16",
            "month": "int8",
            "metric": "string",
            "zone": "string",
            "source_sheet": "string",
        }
    )


def read_series_province() -> pd.DataFrame:
    """Province × year for crashes and victims by severity and zone, 1993–2024."""
    frames: list[pd.DataFrame] = []
    for sheet, (metric, zone) in PROVINCE_SHEETS.items():
        rows = _rows(SERIES_PATH, sheet)
        header = _find_header(rows, "Provincias")
        years = [int(cell) for cell in rows[header][1:] if cell is not None]
        for row in rows[header + 1 :]:
            name = _text(row[0])
            if not name:
                continue
            values = [_number(v) for v in row[1 : 1 + len(years)]]
            if all(v is None for v in values):
                continue
            frames.append(
                pd.DataFrame(
                    {
                        "province": name,
                        "is_total": name.upper() == "TOTAL",
                        "year": years,
                        "metric": metric,
                        "zone": zone,
                        "value": values,
                        "source_sheet": sheet,
                    }
                )
            )
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "province": "string",
            "year": "int16",
            "metric": "string",
            "zone": "string",
            "source_sheet": "string",
        }
    )


def _read_two_level_sheets(
    sheets: dict[str, str], group_field: str, sub_field: str
) -> pd.DataFrame:
    """Sheets with a group header row and a sub-header row (severity × age band, sex × severity).

    Whichever of the two fields is ``severity`` is normalised with :data:`SEVERITY_LABELS`.
    """
    frames: list[pd.DataFrame] = []
    for sheet, zone in sheets.items():
        rows = _rows(SERIES_PATH, sheet)
        header = _find_header(rows, "Años")
        groups = rows[header]
        subs = rows[header + 1]
        columns: list[tuple[str, str]] = []
        current = ""
        for group, sub in zip(groups[1:], subs[1:]):
            if group is not None and _text(group):
                current = _text(group)
            columns.append((current, _text(sub)))
        records: list[dict[str, object]] = []
        for year, *values in _year_rows(rows, header + 1):
            for (group, sub), value in zip(columns, values):
                if not sub:
                    continue
                record = {"year": year, group_field: group, sub_field: sub}
                record["severity"] = SEVERITY_LABELS.get(record["severity"], record["severity"])
                record.update({"zone": zone, "value": _number(value), "source_sheet": sheet})
                records.append(record)
        frames.append(pd.DataFrame.from_records(records))
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "year": "int16",
            group_field: "string",
            sub_field: "string",
            "zone": "string",
            "source_sheet": "string",
        }
    )


def read_series_age() -> pd.DataFrame:
    """Victims by age band × severity × zone, 1993–2024."""
    return _read_two_level_sheets(AGE_SHEETS, "severity", "age_band")


def read_series_sex() -> pd.DataFrame:
    """Victims by sex × severity × zone, 1993–2024."""
    return _read_two_level_sheets(SEX_SHEETS, "sex", "severity")


def read_series_road_users() -> pd.DataFrame:
    """Driver (and passenger) casualties by vehicle type × year × zone, 1993–2024."""
    records: list[dict[str, object]] = []
    for sheet, (population, severity, zone) in ROAD_USER_SHEETS.items():
        rows = _rows(SERIES_PATH, sheet)
        header = _find_header(rows, "Años")
        types = [_text(cell) for cell in rows[header][1:] if cell is not None]
        for year, *values in _year_rows(rows, header):
            for vehicle_type, value in zip(types, values):
                records.append(
                    {
                        "year": year,
                        "population": population,
                        "severity": severity,
                        "vehicle_type": vehicle_type,
                        "is_total": vehicle_type.upper() == "TOTAL",
                        "zone": zone,
                        "value": _number(value),
                        "source_sheet": sheet,
                    }
                )
    out = pd.DataFrame.from_records(records)
    return out.astype(
        {
            "year": "int16",
            "population": "string",
            "severity": "string",
            "vehicle_type": "string",
            "zone": "string",
            "source_sheet": "string",
        }
    )


def read_series_pedestrians() -> pd.DataFrame:
    """Pedestrian victims by severity × zone and by age band × zone, 1993–2024."""
    records: list[dict[str, object]] = []
    for breakdown, sheets in (("severity", PEDESTRIAN_SHEETS), ("age_band", PEDESTRIAN_AGE_SHEETS)):
        for sheet, zone in sheets.items():
            rows = _rows(SERIES_PATH, sheet)
            header = _find_header(rows, "Años")
            labels = [_text(cell) for cell in rows[header][1:] if cell is not None]
            for year, *values in _year_rows(rows, header):
                for label, value in zip(labels, values):
                    category = (
                        SEVERITY_LABELS.get(label, label) if breakdown == "severity" else label
                    )
                    records.append(
                        {
                            "year": year,
                            "breakdown": breakdown,
                            "category": category,
                            "zone": zone,
                            "value": _number(value),
                            "source_sheet": sheet,
                        }
                    )
    out = pd.DataFrame.from_records(records)
    return out.astype(
        {
            "year": "int16",
            "breakdown": "string",
            "category": "string",
            "zone": "string",
            "source_sheet": "string",
        }
    )


# --------------------------------------------------------------------------- 2024 statistical tables

PROVINCE_TABLE_METRICS = (
    "crashes",
    "fatal_crashes",
    "deaths_30d",
    "hospitalised_30d",
    "non_hospitalised_30d",
)


def read_table_1_1() -> pd.DataFrame:
    """TABLA 1.1: 2024 crashes and victims by province and zone."""
    rows = _rows(TABLES_2024_PATH, "TABLA 1.1")
    header = _find_header(rows, "PROVINCIAS")
    zones = ("interurban", "urban", "all")
    frames: list[pd.DataFrame] = []
    for row in rows[header + 2 :]:
        name = _text(row[0])
        if not name:
            continue
        values = [_number(v) for v in row[1:16]]
        if all(v is None for v in values):
            continue
        for zone_index, zone in enumerate(zones):
            block = values[zone_index * 5 : zone_index * 5 + 5]
            frames.append(
                pd.DataFrame(
                    {
                        "province": name,
                        "is_total": name.upper() == "TOTAL",
                        "zone": zone,
                        "metric": list(PROVINCE_TABLE_METRICS),
                        "value": block,
                        "source_sheet": "TABLA 1.1",
                    }
                )
            )
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {"province": "string", "zone": "string", "metric": "string", "source_sheet": "string"}
    )


MONTH_TABLE_METRICS = (
    "crashes",
    "fatal_crashes",
    "victims_30d",
    "deaths_30d",
    "hospitalised_30d",
    "non_hospitalised_30d",
)


def read_table_3_1() -> pd.DataFrame:
    """TABLA 3.1: 2024 crashes and victims by month and zone."""
    rows = _rows(TABLES_2024_PATH, "TABLA 3.1")
    header = _find_header(rows, "MES")
    zones = ("interurban", "urban", "all")
    frames: list[pd.DataFrame] = []
    for row in rows[header + 3 :]:
        name = _text(row[0])
        if not name:
            continue
        month = MONTH_NAMES.index(name) + 1 if name in MONTH_NAMES else None
        values = [_number(v) for v in row[1:19]]
        if all(v is None for v in values):
            continue
        for zone_index, zone in enumerate(zones):
            block = values[zone_index * 6 : zone_index * 6 + 6]
            frames.append(
                pd.DataFrame(
                    {
                        "month_name": name,
                        "month": month,
                        "is_total": name.upper() == "TOTAL",
                        "zone": zone,
                        "metric": list(MONTH_TABLE_METRICS),
                        "value": block,
                        "source_sheet": "TABLA 3.1",
                    }
                )
            )
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "month_name": "string",
            "month": "Int8",
            "zone": "string",
            "metric": "string",
            "source_sheet": "string",
        }
    )


VEHICLE_TABLE_METRICS = ("crashes", "fatal_crashes", "injury_crashes")
VEHICLE_TABLE_YEARS = tuple(range(2020, 2025))
UNIT_ROLES = ("total", "driver", "passenger", "pedestrian")
UNIT_VICTIM_METRICS = (
    "involved",
    "victims_30d",
    "deaths_30d",
    "hospitalised_30d",
    "non_hospitalised_30d",
)


def read_table_2_3(year: int = 2024) -> pd.DataFrame:
    """TABLA 2.3: units (vehicles and pedestrians) involved, by type, zone and crash severity.

    The five workbooks 2020–2024 share the layout: one row per unit type, three zone blocks of three
    severity columns (injury crashes, 30-day fatal crashes, crashes with injured only).
    """
    rows = _rows(tables_raw_path(year), "TABLA 2.3")
    header = _find_header(rows, "TIPO VEHÍCULO")
    zones = ("all", "interurban", "urban")
    frames: list[pd.DataFrame] = []
    for row in rows[header + 2 :]:
        name = _text(row[0])
        if not name:
            continue
        values = [_number(v) for v in row[1:10]]
        if all(v is None for v in values):
            continue
        for zone_index, zone in enumerate(zones):
            block = values[zone_index * 3 : zone_index * 3 + 3]
            frames.append(
                pd.DataFrame(
                    {
                        "year": year,
                        "unit_type": name,
                        "is_total": name.upper() == "TOTAL",
                        "zone": zone,
                        "metric": list(VEHICLE_TABLE_METRICS),
                        "value": block,
                        "source_sheet": "TABLA 2.3",
                    }
                )
            )
        if name.upper() == "TOTAL":
            break
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "year": "int16",
            "unit_type": "string",
            "zone": "string",
            "metric": "string",
            "source_sheet": "string",
        }
    )


def read_units_by_type_all(years: tuple[int, ...] = VEHICLE_TABLE_YEARS) -> pd.DataFrame:
    out = pd.concat([read_table_2_3(year) for year in years], ignore_index=True)
    close_workbooks()
    return out


def read_table_2_2(year: int) -> pd.DataFrame:
    """TABLA 2.2: victims by means of transport, role (driver, passenger, pedestrian) and zone.

    2020–2022 publish one sheet with the interurban and urban blocks side by side; 2023–2024 split it
    into ``2.2.I`` and ``2.2.U``. Each zone block is four role blocks of five columns: units
    involved, victims, deaths, hospitalised and non-hospitalised (30-day definitions).
    """
    path = tables_raw_path(year)
    names = {_sheet_key(name): name for name in _workbook(path).sheetnames}
    sheets = [("2.2", None)] if "2.2" in names else [("2.2.I", "interurban"), ("2.2.U", "urban")]
    frames: list[pd.DataFrame] = []
    for key, zone in sheets:
        rows = _rows(path, names[key])
        header = _find_header(rows, "CLASES DE USUARIOS")
        if zone is None:
            if _text(rows[header][1]).lower() != "vías interurbanas":
                raise ValueError(f"TABLA 2.2 {year}: expected interurban and urban blocks")
            zones: tuple[str, ...] = ("interurban", "urban")
        else:
            zones = (zone,)
        width = len(UNIT_ROLES) * len(UNIT_VICTIM_METRICS)
        for row in rows[header + 1 :]:
            name = _text(row[0])
            if not name or not isinstance(row[1], (int, float)):
                continue
            values = [_number(v) for v in row[1 : 1 + width * len(zones)]]
            for zone_index, block_zone in enumerate(zones):
                for role_index, role in enumerate(UNIT_ROLES):
                    start = zone_index * width + role_index * len(UNIT_VICTIM_METRICS)
                    frames.append(
                        pd.DataFrame(
                            {
                                "year": year,
                                "unit_type": name,
                                "is_total": name.upper() == "TOTAL",
                                "zone": block_zone,
                                "role": role,
                                "metric": list(UNIT_VICTIM_METRICS),
                                "value": values[start : start + len(UNIT_VICTIM_METRICS)],
                                "source_sheet": names[key],
                            }
                        )
                    )
            if name.upper() == "TOTAL":
                break
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "year": "int16",
            "unit_type": "string",
            "zone": "string",
            "role": "string",
            "metric": "string",
            "source_sheet": "string",
        }
    )


def read_victims_by_mode_all(years: tuple[int, ...] = VEHICLE_TABLE_YEARS) -> pd.DataFrame:
    out = pd.concat([read_table_2_2(year) for year in years], ignore_index=True)
    close_workbooks()
    return out


def read_table_8_1_1() -> pd.DataFrame:
    """TABLA 8.1.1: 2024 crashes and victims by number of vehicles involved and zone."""
    rows = _rows(TABLES_2024_PATH, "TABLA 8.1.1")
    header = _find_header(rows, "VEHÍCULOS IMPLICADOS")
    zones = ("all", "interurban", "urban")
    frames: list[pd.DataFrame] = []
    for row in rows[header + 2 :]:
        name = _text(row[0])
        if not name:
            continue
        values = [_number(v) for v in row[1:16]]
        if all(v is None for v in values):
            continue
        for zone_index, zone in enumerate(zones):
            block = values[zone_index * 5 : zone_index * 5 + 5]
            frames.append(
                pd.DataFrame(
                    {
                        "vehicles_involved": name,
                        "is_total": name.upper() == "TOTAL",
                        "zone": zone,
                        "metric": list(PROVINCE_TABLE_METRICS),
                        "value": block,
                        "source_sheet": "TABLA 8.1.1",
                    }
                )
            )
    out = pd.concat(frames, ignore_index=True)
    return out.astype(
        {
            "vehicles_involved": "string",
            "zone": "string",
            "metric": "string",
            "source_sheet": "string",
        }
    )


# --------------------------------------------------------------------------- driver tables 2014–2024

DRIVER_SEX_LABELS = {
    "v": "male",
    "hombre": "male",
    "hombres": "male",
    "m": "female",
    "mujer": "female",
    "mujeres": "female",
    "desconocido": "unknown",
    "se desconoce": "unknown",
}
DRIVER_SEVERITIES = ("deaths_30d", "hospitalised_30d", "non_hospitalised_30d")
DRIVER_SEVERITY_HEADERS = ("fallecidos", "heridos hospitalizados", "heridos no hospitalizados")
DRIVER_TOTAL_LABELS = {"total", "totales"}
DRIVER_ZONES = {"I": "interurban", "U": "urban"}
CHILD_BAND = "0-14"


def _sheet_key(name: object) -> str:
    return _text(name).upper().removeprefix("TABLA ").strip()


def _rows_any(path: Path, sheet_key: str) -> list[tuple]:
    """Rows of the sheet whose name (without the 'TABLA ' prefix) is ``sheet_key``.

    Legacy ``.xls`` workbooks (2014) go through pandas/xlrd, everything else through openpyxl.
    """
    if path.suffix == ".xls":
        book = pd.ExcelFile(path)
        names = {_sheet_key(name): name for name in book.sheet_names}
        frame = book.parse(names[sheet_key], header=None)
        return [
            tuple(None if pd.isna(value) else value for value in row)
            for row in frame.itertuples(index=False)
        ]
    names = {_sheet_key(name): name for name in _workbook(path).sheetnames}
    return _rows(path, names[sheet_key])


def _driver_band(parsed: tuple[int, int | None] | None) -> str:
    if parsed is None:
        return agebands.UNKNOWN
    key = agebands.band_for(parsed[0], parsed[1], agebands.DGT_BANDS)
    if key is None and parsed[1] is not None and parsed[1] <= 14:
        return CHILD_BAND
    if key is None:
        raise ValueError(f"driver table: age {parsed} outside the driver bands")
    return key


def _driver_header(rows: list[tuple]) -> int:
    for index, row in enumerate(rows):
        if len(row) > 2 and _text(row[2]).lower() in DRIVER_TOTAL_LABELS:
            return index
    raise ValueError("driver table: header row with 'Total' in the third column not found")


def _driver_body(rows: list[tuple], first_index: int) -> tuple[list[tuple[str, str, tuple]], tuple]:
    """``[(age_label, sex, row), ...]`` for the data rows, and the published grand-total row.

    Per-age total rows and the grand-total group are skipped; the last total-like row of the sheet
    is returned so callers can verify the parsed sum against the published total.
    """
    body: list[tuple[str, str, tuple]] = []
    grand_total: tuple | None = None
    current: str | None = None
    for row in rows[first_index:]:
        first, second = _text(row[0]), _text(row[1]) if len(row) > 1 else ""
        if first.lower() in DRIVER_TOTAL_LABELS and not second:
            grand_total = row
            continue
        if first.lower() in DRIVER_TOTAL_LABELS:
            current = None  # the grand-total group (2015 onwards): rows by sex, then 'Total'
        elif first:
            current = first
        sex = DRIVER_SEX_LABELS.get(second.lower())
        if current is None and second.lower() == "total":
            grand_total = row
        if current is None or sex is None:
            continue
        body.append((current, sex, row))
    if grand_total is None:
        raise ValueError("driver table: grand total row not found")
    return body, grand_total


def read_table_4_1_1(year: int) -> pd.DataFrame:
    """Table 4.1.1: driver victims by age band, sex and vehicle type, interurban and urban.

    One row per age × sex × vehicle type × severity (killed, hospitalised, not hospitalised, all
    30-day). Sexes are recorded as published (male, female, unknown); totals are not stored, they
    are sums. The parsed rows must reproduce the published grand total of deaths.
    """
    frames: list[pd.DataFrame] = []
    for suffix, zone in DRIVER_ZONES.items():
        sheet = f"4.1.1.{suffix}"
        rows = _rows_any(_driver_table_path(year), sheet)
        header = _driver_header(rows)
        blocks: list[tuple[int, str]] = [
            (column, _text(label))
            for column, label in enumerate(rows[header])
            if column >= 2 and _text(label)
        ]
        sub = [_text(cell).lower() for cell in rows[header + 1]]
        for column, _ in blocks:
            if tuple(sub[column : column + 3]) != DRIVER_SEVERITY_HEADERS:
                raise ValueError(f"table 4.1.1 {year} {zone}: severity headers moved")
        body, grand_total = _driver_body(rows, header + 2)
        records: list[dict[str, object]] = []
        for age_label, sex, row in body:
            parsed = agebands.parse_age_label(age_label)
            for column, vehicle in blocks:
                is_total = vehicle.lower() in DRIVER_TOTAL_LABELS
                for offset, severity in enumerate(DRIVER_SEVERITIES):
                    records.append(
                        {
                            "year": year,
                            "zone": zone,
                            "age_label": age_label,
                            "age_low": None if parsed is None else parsed[0],
                            "age_high": None if parsed is None else parsed[1],
                            "band": _driver_band(parsed),
                            "sex": sex,
                            "vehicle_type": "Total" if is_total else vehicle,
                            "is_total": is_total,
                            "severity": severity,
                            "value": _number(row[column + offset]) or 0.0,
                            "source_sheet": sheet,
                        }
                    )
        frame = pd.DataFrame.from_records(records)
        parsed_deaths = frame[frame.is_total & (frame.severity == "deaths_30d")].value.sum()
        published = _number(grand_total[blocks[0][0]])
        if parsed_deaths != published:
            raise ValueError(
                f"table 4.1.1 {year} {zone}: parsed deaths {parsed_deaths} != published {published}"
            )
        frames.append(frame)
    return _driver_frame(pd.concat(frames, ignore_index=True))


def read_table_4_2(year: int) -> pd.DataFrame:
    """Table 4.2: drivers involved in injury crashes by age band, sex and vehicle type, by zone."""
    frames: list[pd.DataFrame] = []
    for suffix, zone in DRIVER_ZONES.items():
        sheet = f"4.2.{suffix}"
        rows = _rows_any(_driver_table_path(year), sheet)
        header = _driver_header(rows)
        columns = [
            (column, _text(label))
            for column, label in enumerate(rows[header])
            if column >= 2 and _text(label)
        ]
        body, grand_total = _driver_body(rows, header + 1)
        records: list[dict[str, object]] = []
        for age_label, sex, row in body:
            parsed = agebands.parse_age_label(age_label)
            for column, vehicle in columns:
                is_total = vehicle.lower() in DRIVER_TOTAL_LABELS
                records.append(
                    {
                        "year": year,
                        "zone": zone,
                        "age_label": age_label,
                        "age_low": None if parsed is None else parsed[0],
                        "age_high": None if parsed is None else parsed[1],
                        "band": _driver_band(parsed),
                        "sex": sex,
                        "vehicle_type": "Total" if is_total else vehicle,
                        "is_total": is_total,
                        "value": _number(row[column]) or 0.0,
                        "source_sheet": sheet,
                    }
                )
        frame = pd.DataFrame.from_records(records)
        parsed_total = frame[frame.is_total].value.sum()
        published = _number(grand_total[columns[0][0]])
        if parsed_total != published:
            raise ValueError(
                f"table 4.2 {year} {zone}: parsed total {parsed_total} != published {published}"
            )
        frames.append(frame)
    return _driver_frame(pd.concat(frames, ignore_index=True))


def _driver_table_path(year: int) -> Path:
    return tables_raw_path(year, 4)


def _driver_frame(frame: pd.DataFrame) -> pd.DataFrame:
    types = {
        "year": "int16",
        "zone": "string",
        "age_label": "string",
        "age_low": "Int16",
        "age_high": "Int16",
        "band": "string",
        "sex": "string",
        "vehicle_type": "string",
        "source_sheet": "string",
    }
    if "severity" in frame:
        types["severity"] = "string"
    return frame.astype(types)


def read_driver_victims_all(years: tuple[int, ...] = TABLE_YEARS) -> pd.DataFrame:
    out = pd.concat([read_table_4_1_1(year) for year in years], ignore_index=True)
    close_workbooks()
    return out


def read_drivers_involved_all(years: tuple[int, ...] = TABLE_YEARS) -> pd.DataFrame:
    out = pd.concat([read_table_4_2(year) for year in years], ignore_index=True)
    close_workbooks()
    return out


# --------------------------------------------------------------------------- driver infractions 6.1

INFRACTION_BLOCKS = ("speed", "driver", "door", "lighting", "load", "summary")
INFRACTION_VEHICLES = {
    "total": "total",
    "bicicleta": "bicycle",
    "vmp": "vmp",
    "ciclomotor": "moped",
    "motocicleta": "motorcycle",
    "turismo": "car",
    "furgoneta": "van",
    "camión hasta": "light_truck",
    "camión más": "heavy_truck",
    "autobús": "bus",
    "otro vehículo": "other",
    "se desconoce": "unknown",
    "sin dato": "unknown",
    "peatón": None,  # 2015 only, all zeros
}
# Item labels that identify their block; the shared ones (none, unknown, total) take the current
# block, and a repeated item advances to the next block.
INFRACTION_ITEMS = {
    "infracción de velocidad": ("speed", "speed_infraction"),
    "marcha lenta": ("speed", "too_slow"),
    "ninguna infracción de velocidad": ("speed", "none"),
    "no respetar señal de stop": ("driver", "stop_sign"),
    "no respetar paso para peatones": ("driver", "pedestrian_crossing"),
    "no respetar otra regulación": ("driver", "other_priority"),
    "circular en sentido contrario": ("driver", "wrong_way"),
    "invadir parcialmente": ("driver", "partial_wrong_side"),
    "adelantar antirreglamentariamente": ("driver", "overtaking"),
    "no mantener el intervalo": ("driver", "safety_distance"),
    "otra infracción": ("driver", "other_infraction"),
    "apertura de puertas": ("door", "door_opening"),
    "incorrecta utilización del alumbrado": ("lighting", "lighting"),
    "exceso, mal acondicionamiento": ("load", "load"),
    "alguna infracción": ("summary", "any"),
}
INFRACTION_SHARED = {
    "ninguna infracción": "none",
    "se desconoce": "unknown",
    "se ignora": "unknown",
    "total": "total",
}
INFRACTION_PREFIXES = (
    "infracciones de velocidad",
    "infracciones del conductor",
    "infracciones de apertura de puerta",
    "infracciones de alumbrado",
    "infracciones de carga del vehículo",
    "resumen de infracciones",
)
INFRACTION_ITEM_LABELS = {
    "speed_infraction": "Speed infraction",
    "too_slow": "Driving too slowly, obstructing traffic",
    "stop_sign": "Failing to stop at a stop sign",
    "pedestrian_crossing": "Failing to respect a pedestrian crossing",
    "other_priority": "Other priority infraction",
    "wrong_way": "Driving against the flow or where prohibited",
    "partial_wrong_side": "Partly invading the opposite lane",
    "overtaking": "Illegal overtaking",
    "safety_distance": "Not keeping a safe distance",
    "other_infraction": "Other infraction",
    "door_opening": "Opening a door without care",
    "lighting": "Incorrect use of lights",
    "load": "Excess, badly secured or shed load",
    "any": "Some infraction",
    "none": "No infraction",
    "unknown": "Unknown",
    "total": "All drivers",
}


def _infraction_key(label: str) -> str:
    text = " ".join(label.lower().split())
    for prefix in INFRACTION_PREFIXES:
        if text.startswith(prefix + " "):
            text = text[len(prefix) + 1 :]
            break
    return text


def _classify_infraction(text: str, current: str, seen: set[tuple[str, str]]) -> tuple[str, str]:
    """(block, item) for a row label given the current block and the items already recorded."""
    for pattern, (block, item) in INFRACTION_ITEMS.items():
        if text.startswith(pattern):
            return block, item
    if text.startswith("total conductores") or text.startswith("total resumen"):
        return "summary", "total"
    for pattern, item in INFRACTION_SHARED.items():
        if text == pattern or text.startswith(pattern + " "):
            block = current
            if (block, item) in seen:
                block = INFRACTION_BLOCKS[min(INFRACTION_BLOCKS.index(block) + 1, 5)]
            return block, item
    raise ValueError(f"table 6.1: unknown row label {text!r}")


def read_table_6_1(year: int, zone: str) -> pd.DataFrame:
    """TABLA 6.1.I / 6.1.U: drivers involved by vehicle type × infraction, one year and zone.

    Long format: ``year``, ``zone``, ``block`` (speed, driver, door, lighting, load, summary),
    ``item``, ``vehicle_group``, ``value``. The block totals are the drivers involved. Rows are
    located by label, so the four layouts of 2014–2024 parse the same way.
    """
    zone_code = {"interurban": "I", "urban": "U"}[zone]
    rows = _rows_any(tables_raw_path(year, 6), f"6.1.{zone_code}")
    header_index = next(
        i for i, row in enumerate(rows) if any("bicicleta" in _text(c).lower() for c in row)
    )
    columns: dict[int, str] = {}
    for index, cell in enumerate(rows[header_index]):
        text = _text(cell).lower()
        if not text:
            continue
        for pattern, group in INFRACTION_VEHICLES.items():
            if text.startswith(pattern):
                if group is not None:
                    columns[index] = group
                break
    if "total" not in columns.values() or "car" not in columns.values():
        raise ValueError(f"table 6.1 {year} {zone}: vehicle columns not found")
    first_value = min(columns)
    records: list[dict[str, object]] = []
    current = "speed"
    seen: set[tuple[str, str]] = set()
    for row in rows[header_index + 1 :]:
        label_cells = [_text(c) for c in row[:first_value] if isinstance(c, str) and _text(c)]
        values = [_number(row[i]) if i < len(row) else None for i in columns]
        if not label_cells or all(v is None for v in values):
            continue
        text = _infraction_key(label_cells[-1])
        if text.startswith("debido a"):
            break
        block, item = _classify_infraction(text, current, seen)
        current = block
        seen.add((block, item))
        for (index, group), value in zip(columns.items(), values):
            records.append(
                {
                    "year": year,
                    "zone": zone,
                    "block": block,
                    "item": item,
                    "vehicle_group": group,
                    "value": 0.0 if value is None else float(value),
                }
            )
        if item == "total" and block != "summary":
            current = INFRACTION_BLOCKS[INFRACTION_BLOCKS.index(block) + 1]
    out = pd.DataFrame.from_records(records)
    keys = ["year", "zone", "block", "item", "vehicle_group"]
    if out.duplicated(keys).any():
        raise ValueError(f"table 6.1 {year} {zone}: a row was assigned to a block twice")
    return out.astype(
        {
            "year": "int16",
            "zone": "string",
            "block": "string",
            "item": "string",
            "vehicle_group": "string",
        }
    )


def read_driver_infractions_all(years: tuple[int, ...] = TABLE_YEARS) -> pd.DataFrame:
    frames = [read_table_6_1(year, zone) for year in years for zone in ("interurban", "urban")]
    out = pd.concat(frames, ignore_index=True)
    close_workbooks()
    return out


# --------------------------------------------------------------------------- interim layer

TABLE_BUILDERS = {
    "series_annual": read_series_annual,
    "series_monthly": read_series_monthly,
    "series_province": read_series_province,
    "series_age": read_series_age,
    "series_sex": read_series_sex,
    "series_road_users": read_series_road_users,
    "series_pedestrians": read_series_pedestrians,
    "tables_2024_province": read_table_1_1,
    "tables_2024_month": read_table_3_1,
    "tables_units_by_type": read_units_by_type_all,
    "tables_victims_by_mode": read_victims_by_mode_all,
    "tables_2024_vehicles_involved": read_table_8_1_1,
    "tables_driver_victims": read_driver_victims_all,
    "tables_drivers_involved": read_drivers_involved_all,
    "tables_driver_infractions": read_driver_infractions_all,
}


def interim_path(name: str) -> Path:
    return INTERIM_DATA_DIR / f"{name}.parquet"


def build_tables(force: bool = False) -> list[Path]:
    """Write every published-table frame to ``data/interim`` as Parquet."""
    written: list[Path] = []
    INTERIM_DATA_DIR.mkdir(parents=True, exist_ok=True)
    for name, builder in TABLE_BUILDERS.items():
        target = interim_path(name)
        if target.exists() and not force:
            log.info("tables %s: exists, skipping", name)
            written.append(target)
            continue
        frame = builder()
        frame.to_parquet(target, index=False)
        log.info("tables %s: %s rows -> %s", name, f"{len(frame):,}", target.name)
        written.append(target)
    close_workbooks()
    return written


def read_table(name: str) -> pd.DataFrame:
    return pd.read_parquet(interim_path(name))
